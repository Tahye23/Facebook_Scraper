"""Generation Word optionnelle du rapport Mauritanie 24h.

IMPORTANT: ce module ne doit JAMAIS etre importe au demarrage du worker.
Il est charge en lazy-import uniquement apres le scraping CSV 24h, et toute
erreur est ignoree par l'appelant pour ne pas casser le job.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path


def _int_value(value) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def _author_label(author: str | None, source: str | None = None) -> str:
    text = str(author or "").strip().lstrip("@")
    if text:
        return text
    raw = str(source or "").strip()
    return raw[:48] if raw else "مصدر غير معروف"


def build_mauritanie_24h_docx(
    scrape_id: str,
    gemini_report: dict | None,
    videos_payload: list[dict] | None,
    output_dir: Path,
) -> str | None:
    """Ecrit mauritanie_24h_{scrape_id}.docx et retourne le chemin, ou None."""
    videos = [v for v in (videos_payload or []) if isinstance(v, dict)]
    if not videos:
        return None

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except Exception:
        return None

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"mauritanie_24h_{scrape_id}.docx"

    title = str((gemini_report or {}).get("report_title_ar") or "موريتانيا في وسائل التواصل الاجتماعي")
    overview = str((gemini_report or {}).get("overview_ar") or "")
    conclusion = str((gemini_report or {}).get("conclusion_ar") or "")

    total_views = sum(_int_value(v.get("views")) for v in videos)
    total_likes = sum(_int_value(v.get("likes")) for v in videos)
    total_comments = sum(_int_value(v.get("comments")) for v in videos)

    doc = Document()
    h = doc.add_heading(title, 0)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)

    p = doc.add_paragraph("تقرير تحليلي للمحتوى المستخرج")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d = doc.add_paragraph(datetime.now(tz=timezone.utc).strftime("%Y-%m-%d"))
    d.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_heading("أولاً: نظرة عامة على النشاط الرقمي", level=1)
    if overview:
        op = doc.add_paragraph(overview)
        op.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for line in (
        f"إجمالي المنشورات: {len(videos)}",
        f"إجمالي المشاهدات: {total_views:,}",
        f"إجمالي الإعجابات: {total_likes:,}",
        f"إجمالي التعليقات: {total_comments:,}",
    ):
        lp = doc.add_paragraph(line)
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_heading("ثانياً: تفاصيل المنشورات", level=1)
    for item in videos:
        author = _author_label(item.get("author"), item.get("source") or item.get("post_url"))
        text = str(item.get("description") or "").strip() or "لا يوجد نص"
        hp = doc.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(f"@{author}")
        run.bold = True
        run.font.size = Pt(11)
        tp = doc.add_paragraph(text)
        tp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        mp = doc.add_paragraph(
            f"views={_int_value(item.get('views'))} likes={_int_value(item.get('likes'))} "
            f"comments={_int_value(item.get('comments'))} shares={_int_value(item.get('shares'))}"
        )
        mp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        url = str(item.get("post_url") or "").strip()
        if url:
            up = doc.add_paragraph(url)
            up.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    if conclusion:
        doc.add_heading("الخلاصة", level=1)
        cp = doc.add_paragraph(conclusion)
        cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.save(str(out_path))
    return str(out_path)
